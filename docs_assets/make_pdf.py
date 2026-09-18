"""Build Darukaa.Earth solution DOCX. Run: python docs_assets/make_pdf.py"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO = "https://github.com/Palakagarwal28/Darukaa.earth-hackathon"
A = "docs_assets/"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)

def h1(t): doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)
def p(t, bold=False, italic=False):
    pa = doc.add_paragraph()
    r = pa.add_run(t); r.bold = bold; r.italic = italic
    return pa
def bullets(items):
    for it in items: doc.add_paragraph(it, style="List Bullet")
def table(rows):
    tb = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tb.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tb.cell(i, j); c.text = val
            for par in c.paragraphs:
                for r in par.runs: r.font.size = Pt(10)
    doc.add_paragraph()
def img(path, width=6.0, caption=""):
    doc.add_picture(path, width=Inches(width))
    if caption:
        pa = doc.add_paragraph(); pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = pa.add_run(caption); r.italic = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

# ---- cover ----
t = doc.add_heading("Darukaa.Earth: AI Biodiversity Intelligence Chatbot", level=0)
p("Hackathon solution document — AI environmental scientist, not a chatbot.", italic=True)
p("GitHub repository: " + REPO, bold=True)
p("The repository contains the complete project: backend, knowledge base, frontend, tests, CI.")
p("Runs fully offline — no deployment, credentials, or API keys required.", italic=True)

h1("1. Objective")
p("Build an AI-powered conversational system that maintains a structured knowledge base of "
  "biodiversity and environmental metrics, understands user queries about ecosystems, land and "
  "climate, generates actionable non-obvious recommendations, and supports every recommendation "
  "with scientific reasoning and evidence.")

h1("2. Tech stack")
table([
    ["Layer", "Technology", "Why"],
    ["API backend", "FastAPI (Python) + Uvicorn", "Typed endpoints, auto OpenAPI docs at /docs"],
    ["Retrieval (RAG)", "BM25 + agro synonym map (default); TF-IDF / dense optional", "7/7 benchmark, offline, explainable scores"],
    ["Knowledge store", "Versioned JSON (32 evidence cards) + schema", "Auditable, migratable to pgvector/Chroma"],
    ["Reasoning", "Deterministic multi-metric engine (no LLM in loop)", "No hallucinated advice, reproducible"],
    ["Conversation", "Slot-filling + SQLite multi-turn memory", "Clarifying Qs, context awareness"],
    ["Frontend", "Vanilla HTML/CSS/JS, zero build step", "Runs from the same server, 3 tabs"],
    ["Enrichment (optional)", "Open-Meteo (keyless) + OpenAI-compatible LLM rephrase", "Graceful offline fallback"],
    ["Quality", "pytest (12 tests), GitHub Actions CI", "Green on every push"],
    ["Docs deliverable", "This document (DOCX->PDF via MS Word)", "Submission artifact"],
])

h1("3. System architecture")
img(A + "architecture.png", 6.0, "Figure 1 — end-to-end architecture: deterministic engine + RAG evidence.")
table([
    ["Stage", "Module", "What it does"],
    ["1. Input", "Chat / Assess endpoints", "Text, structured JSON, geo lat/lon (bonus)"],
    ["2. Conversation", "src/conversation.py", "Extracts 5-domain metrics; asks clarifying Qs if soc/rainfall/land_use missing; SQLite memory"],
    ["3. Enrichment", "src/geo.py, src/integrations.py", "Offline region inference; optional Open-Meteo fill (flagged, user data wins)"],
    ["4. Retrieval", "src/retriever.py, src/bm25.py", "BM25 over 32 cards; returns kb_id + score + matched terms"],
    ["5. Reasoning", "src/reasoner.py", "Interaction diagnosis, confidence, sequencing, trade-offs, targets, measure-next"],
    ["6. Response", "UI + JSON API", "Phased plan + recs + provenance [KB-ID] + Markdown export"],
])

h1("4. Knowledge system (critical requirement)")
bullets([
    "32 evidence cards covering soil pH/SOC/moisture, land use/cover, species richness/habitat diversity, "
    "temperature/rainfall, pollution/deforestation — plus measurement protocols (eDNA KB26, soil card KB27).",
    "Sources: FAO, IPCC AR6, IPBES (2016/2019/2023), UNCCD, Ramsar, CBD + peer-reviewed meta-analyses.",
    "Each card: intervention, mechanism, metrics improved, effect size, time horizon, confidence, source, "
    "region fit, conditions, contraindications.",
    "Validation: python -m knowledge.build_index (IDs, fields, coverage, retrieval smoke test).",
])
img(A + "benchmark.png", 5.0, "Figure 2 — retrieval benchmark: BM25+synonyms 7/7 vs TF-IDF 5/7.")
p("The two adversarial farmer paraphrases ('white crust on field, salty' -> KB16 sodic reclamation; "
  "'field floods every monsoon' -> KB28 drainage) miss on TF-IDF and hit on BM25 — the measured "
  "justification for the default backend. Method + full table: knowledge/retrieval_benchmark.md.",
  italic=True)

h1("5. Conversational intelligence")
bullets([
    "Vague input ('biodiversity is declining') triggers targeted clarifying questions, never a guess.",
    "Multi-turn memory per session id (SQLite); JSON overrides memory; new text merges in.",
    "measure_next ranks the 3 missing inputs by decision impact, driving follow-ups.",
    "Inputs: free text, structured JSON (/assess), geo-coordinates (region inference + climate enrich).",
])
img(A + "ui_chat.png", 6.0, "Figure 3 — Chat Scientist tab running the semi-arid wheat demo (?tab=chat&demo=1).")

h1("6. Evidence-backed, multi-metric recommendations")
p("Every recommendation carries WHAT / WHY (mechanism) / impacted metrics / expected effect / "
  "time horizon / confidence + score / reference [KB-ID]. Example (SOC 0.3%, low rain, monoculture wheat):")
bullets([
    "Legume intercrop -> SOC +15-25% in 2-3 yrs, pollinators +30-50% (FAO; Poeplau & Don 2015) [KB01].",
    "Contour bunds + farm pond -> moisture +25-40% (UNCCD; IPCC AR6) [KB04].",
    "Diagnosis chains variables (low SOC x drought x monoculture x pH), never single-variable answers.",
    "Sequenced plan: chemistry first, water early, biology over seasons — plus trade-offs and .md export.",
])
img(A + "ui_assess.png", 6.0, "Figure 4 — Structured Assess tab: phased plan, targets, recs, provenance (?tab=assess&demo=1).")

h1("7. Knowledge grounding surface (for judges)")
bullets([
    "GET /knowledge/search shows scores + matched terms for any query.",
    "GET /provenance?ids= returns the exact cited cards.",
    "12 pytest tests + green GitHub Actions CI on every push.",
])
img(A + "ui_kb.png", 6.0, "Figure 5 — Knowledge explorer tracing 'white crust salty field' to KB16.")

h1("8. API reference")
table([
    ["Endpoint", "Purpose"],
    ["POST /chat", "Text + memory; clarifying questions when incomplete"],
    ["POST /assess", "Structured JSON site -> full scientific report"],
    ["GET /knowledge/search", "Retrieval trace (scores, matched terms)"],
    ["GET /provenance?ids=", "Exact evidence cards cited"],
    ["GET /sessions/{sid}", "Conversation history"],
    ["GET /health, /integrations/status, /docs", "Status + interactive API docs"],
])

h1("9. Evaluation summary")
table([
    ["Metric", "Result", "Where to verify"],
    ["Retrieval benchmark", "BM25 7/7, TF-IDF 5/7", "knowledge/retrieval_benchmark.md"],
    ["Retrieval latency (median)", "BM25 2.3 ms; assess() 2.7 ms; HTTP 13.4 ms", "Rerun snippet in ARCHITECTURE.md"],
    ["Tests", "12/12 pass", "python -m pytest tests -q; CI badge"],
    ["Variables combined", "6 (wheat demo)", "POST /assess variables_count"],
    ["Offline", "Identical results, no network", "Disconnect + rerun demo"],
])

h1("10. Local setup (2 minutes)")
bullets([
    "pip install -r requirements.txt",
    "python -m pytest tests -q",
    "python -m uvicorn api.main:app --host 127.0.0.1 --port 8000  (or double-click run_local.ps1)",
    "Open http://127.0.0.1:8000  |  shareable demos: ?tab=chat&demo=1  ?tab=assess&demo=1  ?tab=kb&demo=1&q=...",
])
p("GitHub: " + REPO + "  |  No deployment, keys, or credentials required.", bold=True)

doc.save("Darukaa_Earth_Solution.docx")
print("docx done")
